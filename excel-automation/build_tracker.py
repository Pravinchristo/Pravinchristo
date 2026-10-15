"""Build the automated Recruitment Tracker workbook.

Usage:
    python build_tracker.py                 # with demo data
    python build_tracker.py --blank         # empty template (one example row)
"""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BLANK = "--blank" in sys.argv
OUT = "Recruitment_Tracker_Template.xlsx" if BLANK else "Recruitment_Tracker.xlsx"
MAX_ROWS = 500          # candidate rows pre-wired with formulas
MAX_REQS = 100          # requisition rows pre-wired with formulas
CL, RL = MAX_ROWS + 1, MAX_REQS + 1

FONT = "Arial"
NAVY, TEAL, GREY = "1F3864", "2E75B6", "F2F2F2"
INPUT_FILL = PatternFill("solid", fgColor="FFF9E5")   # cream = type here
AUTO_FILL = PatternFill("solid", fgColor="EAF1FB")    # light blue = formula
HDR_FILL = PatternFill("solid", fgColor=NAVY)
AUTO_HDR_FILL = PatternFill("solid", fgColor=TEAL)
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


wb = Workbook()

# ------------------------------------------------------------------ Lists
ls = wb.active
ls.title = "Lists"
STAGES = [  # stage, order, SLA days, group, next action
    ("Sourced", 1, 3, "Active", "Reach out & share JD"),
    ("Screening", 2, 3, "Active", "Complete phone screen"),
    ("Tech Round 1", 3, 5, "Active", "Schedule / collect Tech R1 feedback"),
    ("Tech Round 2", 4, 5, "Active", "Schedule / collect Tech R2 feedback"),
    ("HR Round", 5, 3, "Active", "HR discussion & salary alignment"),
    ("Offered", 6, 5, "Active", "Follow up on offer acceptance"),
    ("Offer Accepted", 7, 30, "Active", "Pre-joining engagement & docs"),
    ("Joined", 8, 0, "Hired", "Onboarding complete"),
    ("Offer Declined", 6, 0, "Lost", "Close & capture decline reason"),
    ("Rejected", 0, 0, "Lost", "Send regret email"),
    ("Withdrawn", 0, 0, "Lost", "Close & capture reason"),
]
SOURCES = ["LinkedIn", "Naukri", "Referral", "Career Page", "Indeed",
           "Agency", "Job Fair", "Internal"]
RECRUITERS = ["Pravin Christo", "Recruiter 2", "Recruiter 3"]
DEPTS = ["QA", "Engineering", "DevOps", "Data", "Cloud"]
LOCS = ["Chennai", "Bangalore", "Hyderabad", "Pune", "Remote"]
PRIOS = ["High", "Medium", "Low"]
DROP_STAGES = [s[0] for s in STAGES if s[3] == "Active"]

ls["A1"] = "Lists & Settings — edit these to customise every dropdown and rule"
ls["A1"].font = font(bold=True, size=12, color=NAVY)
heads = ["Stage", "Stage Order", "SLA (days in stage)", "Group", "Next Action"]
for c, h in enumerate(heads, 1):
    ls.cell(row=3, column=c, value=h)
for r, row in enumerate(STAGES, 4):
    for c, v in enumerate(row, 1):
        ls.cell(row=r, column=c, value=v)
STAGE_LAST = 3 + len(STAGES)
ST_RNG = f"Lists!$A$4:$A${STAGE_LAST}"
for col, title, items in [("G", "Sources", SOURCES), ("H", "Recruiters", RECRUITERS),
                          ("I", "Departments", DEPTS), ("J", "Locations", LOCS),
                          ("K", "Priority", PRIOS), ("L", "Drop-off Stage", DROP_STAGES)]:
    ls[f"{col}3"] = title
    for i, v in enumerate(items, 4):
        ls[f"{col}{i}"] = v
for c in range(1, 13):
    cell = ls.cell(row=3, column=c)
    if cell.value:
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
for row in ls.iter_rows(min_row=4, max_row=30, max_col=12):
    for cell in row:
        if cell.value is not None:
            cell.font = font(color="0000FF")
            cell.fill = INPUT_FILL
            cell.border = BORDER
ls["A17"] = ("Stage Order: position in the hiring funnel (0 = not part of funnel). "
             "Group: Active / Hired / Lost drives all status logic. SLA = max days "
             "a candidate should sit in that stage before being flagged.")
ls["A17"].font = font(italic=True, color="595959", size=9)
widths = {"A": 16, "B": 12, "C": 18, "D": 10, "E": 36, "F": 3, "G": 14,
          "H": 18, "I": 14, "J": 13, "K": 10, "L": 16}
for k, v in widths.items():
    ls.column_dimensions[k].width = v
# Extend dropdown ranges to row 30 so the user can add more items
SRC_RNG, REC_RNG, DEP_RNG = "Lists!$G$4:$G$30", "Lists!$H$4:$H$30", "Lists!$I$4:$I$30"
LOC_RNG, PRI_RNG, DROP_RNG = "Lists!$J$4:$J$30", "Lists!$K$4:$K$30", "Lists!$L$4:$L$30"

# ------------------------------------------------------------ Requisitions
rq = wb.create_sheet("Requisitions", 0)
REQ_COLS = [  # header, kind, width, formula(template with {r})
    ("Req ID", "in", 10, None),
    ("Role", "in", 26, None),
    ("Department", "in", 13, None),
    ("Hiring Manager", "in", 18, None),
    ("Location", "in", 12, None),
    ("Openings", "in", 10, None),
    ("Priority", "in", 9, None),
    ("Recruiter", "in", 16, None),
    ("Open Date", "in", 12, None),
    ("Target Close Date", "in", 12, None),
    ("Total Candidates", "auto", 11, '=IF($A{r}="","",COUNTIF(Candidates!$E$2:$E${CL},$A{r}))'),
    ("Active Pipeline", "auto", 10, '=IF($A{r}="","",COUNTIFS(Candidates!$E$2:$E${CL},$A{r},Candidates!$M$2:$M${CL},"Active"))'),
    ("Offers Made", "auto", 9, '=IF($A{r}="","",COUNTIFS(Candidates!$E$2:$E${CL},$A{r},Candidates!$T$2:$T${CL},"<>"))'),
    ("Joined", "auto", 8, '=IF($A{r}="","",COUNTIFS(Candidates!$E$2:$E${CL},$A{r},Candidates!$K$2:$K${CL},"Joined"))'),
    ("Openings Left", "auto", 9, '=IF($A{r}="","",MAX($F{r}-$N{r},0))'),
    ("Fill %", "auto", 8, '=IF(OR($A{r}="",N($F{r})=0),"",$N{r}/$F{r})'),
    ("Days Open", "auto", 8, '=IF(OR($A{r}="",$I{r}=""),"",TODAY()-$I{r})'),
    ("Days to Target", "auto", 9, '=IF(OR($A{r}="",$J{r}=""),"",$J{r}-TODAY())'),
    ("Req Status", "auto", 11, '=IF($A{r}="","",IF($O{r}=0,"Filled",IF(AND($J{r}<>"",TODAY()>$J{r}),"Overdue",IF($L{r}=0,"No Pipeline","Open"))))'),
]
# ------------------------------------------------------------ Candidates
cd = wb.create_sheet("Candidates", 1)
STG_LOOKUP = lambda col: f'INDEX(Lists!${col}$4:${col}${STAGE_LAST},MATCH($K{{r}},{ST_RNG},0))'
CAND_COLS = [
    ("Candidate ID", "auto", 11, '=IF($B{r}="","","C-"&TEXT(ROW()-1,"0000"))'),
    ("Candidate Name", "in", 20, None),
    ("Email", "in", 24, None),
    ("Phone", "in", 13, None),
    ("Req ID", "in", 9, None),
    ("Role", "auto", 24, '=IF($E{r}="","",IFERROR(INDEX(Requisitions!$B$2:$B${RL},MATCH($E{r},Requisitions!$A$2:$A${RL},0)),"Unknown Req"))'),
    ("Hiring Manager", "auto", 16, '=IF($E{r}="","",IFERROR(INDEX(Requisitions!$D$2:$D${RL},MATCH($E{r},Requisitions!$A$2:$A${RL},0)),""))'),
    ("Source", "in", 12, None),
    ("Recruiter", "in", 15, None),
    ("Applied Date", "in", 11, None),
    ("Current Stage", "in", 14, None),
    ("Stage Updated On", "in", 11, None),
    ("Pipeline Status", "auto", 10, '=IF($K{r}="","",IFERROR(' + STG_LOOKUP("D") + ',"Check stage"))'),
    ("Dropped At Stage", "in", 13, None),
    ("Exp (yrs)", "in", 8, None),
    ("Current CTC (LPA)", "in", 10, None),
    ("Expected CTC (LPA)", "in", 10, None),
    ("Notice (days)", "in", 8, None),
    ("Interview Rating (1-5)", "in", 9, None),
    ("Offer Date", "in", 11, None),
    ("Joining Date", "in", 11, None),
    ("Hike %", "auto", 8, '=IF(OR(N($P{r})=0,$Q{r}=""),"",$Q{r}/$P{r}-1)'),
    ("Days in Pipeline", "auto", 9, '=IF($J{r}="","",IF($M{r}="Active",TODAY(),IF($U{r}<>"",$U{r},IF($L{r}<>"",$L{r},TODAY())))-$J{r})'),
    ("Days in Stage", "auto", 8, '=IF(OR($L{r}="",$M{r}<>"Active"),"",TODAY()-$L{r})'),
    ("Stage SLA (days)", "auto", 8, '=IF($M{r}<>"Active","",' + STG_LOOKUP("C") + ')'),
    ("SLA Status", "auto", 10, '=IF($B{r}="","",IF($M{r}<>"Active","Closed",IF($X{r}="","",IF($X{r}>$Y{r},"BREACH","On Track"))))'),
    ("Next Action", "auto", 30, '=IF($K{r}="","",IFERROR(' + STG_LOOKUP("E") + ',""))'),
    ("Expected Joining", "auto", 11, '=IF($U{r}<>"",$U{r},IF(OR($T{r}="",$M{r}<>"Active"),"",$T{r}+N($R{r})))'),
    ("Time to Offer (days)", "auto", 9, '=IF(OR($T{r}="",$J{r}=""),"",$T{r}-$J{r})'),
    ("Time to Hire (days)", "auto", 9, '=IF(AND($K{r}="Joined",$U{r}<>"",$J{r}<>""),$U{r}-$J{r},"")'),
    ("Furthest Stage #", "auto", 9, '=IF($K{r}="","",IF(OR($K{r}="Rejected",$K{r}="Withdrawn"),IFERROR(INDEX(Lists!$B$4:$B${SL},MATCH($N{r},{ST},0)),1),IFERROR(' + STG_LOOKUP("B") + ',0)))'),
]


def build_table(ws, cols, nrows, freeze):
    for c, (h, kind, w, _) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HDR_FILL if kind == "in" else AUTO_HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 42
    for r in range(2, nrows + 2):
        for c, (h, kind, w, f) in enumerate(cols, 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if kind == "auto":
                cell.value = f.format(r=r, CL=CL, RL=RL, SL=STAGE_LAST, ST=ST_RNG)
                cell.fill = AUTO_FILL
                cell.font = font()
            else:
                cell.font = font(color="0000FF")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{nrows + 1}"


build_table(rq, REQ_COLS, MAX_REQS, "C2")
build_table(cd, CAND_COLS, MAX_ROWS, "C2")
for c in range(1, len(CAND_COLS) + 1):
    cd.cell(row=1, column=c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

# number formats
DATE_FMT = "dd-mmm-yy"
for col in "IJ":
    for r in range(2, RL + 1):
        rq[f"{col}{r}"].number_format = DATE_FMT
for r in range(2, RL + 1):
    rq[f"P{r}"].number_format = "0%"
for col in ["J", "L", "T", "U", "AB"]:
    for r in range(2, CL + 1):
        cd[f"{col}{r}"].number_format = DATE_FMT
for col in ["P", "Q"]:
    for r in range(2, CL + 1):
        cd[f"{col}{r}"].number_format = "0.0"
for r in range(2, CL + 1):
    cd[f"V{r}"].number_format = "0%"

# ---------------------------------------------------------- data validation
def dv(ws, formula, rng, title, prompt, allow="list"):
    v = DataValidation(type=allow, formula1=formula, allow_blank=True)
    v.promptTitle, v.prompt = title, prompt
    v.errorTitle, v.error = "Invalid entry", "Please pick a value from the list / enter a valid value."
    v.showErrorMessage = True
    v.showInputMessage = True
    ws.add_data_validation(v)
    v.add(rng)


dv(cd, f"=Requisitions!$A$2:$A${RL}", f"E2:E{CL}", "Req ID", "Pick the requisition this candidate applied for.")
dv(cd, f"={SRC_RNG}", f"H2:H{CL}", "Source", "Where did this candidate come from?")
dv(cd, f"={REC_RNG}", f"I2:I{CL}", "Recruiter", "Owner of this candidate.")
dv(cd, f"={ST_RNG}", f"K2:K{CL}", "Stage", "Update this + 'Stage Updated On' whenever the candidate moves.")
dv(cd, f"={DROP_RNG}", f"N2:N{CL}", "Dropped At", "Only for Rejected / Withdrawn: the stage they dropped out at.")
for col in ["J", "L", "T", "U"]:
    dv(cd, "DATE(2000,1,1)", f"{col}2:{col}{CL}", "Date", "Enter a date, e.g. 15-Oct-26", allow="date")
    cd.data_validations.dataValidation[-1].operator = "greaterThan"
v = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
v.error, v.showErrorMessage = "Rating must be 1 to 5", True
cd.add_data_validation(v); v.add(f"S2:S{CL}")
dv(rq, f"={DEP_RNG}", f"C2:C{RL}", "Department", "")
dv(rq, f"={LOC_RNG}", f"E2:E{RL}", "Location", "")
dv(rq, f"={PRI_RNG}", f"G2:G{RL}", "Priority", "")
dv(rq, f"={REC_RNG}", f"H2:H{RL}", "Recruiter", "")
for col in "IJ":
    dv(rq, "DATE(2000,1,1)", f"{col}2:{col}{RL}", "Date", "Enter a date", allow="date")
    rq.data_validations.dataValidation[-1].operator = "greaterThan"
v = DataValidation(type="custom", formula1=f"COUNTIF($A$2:$A${RL},A2)=1", allow_blank=True)
v.error, v.showErrorMessage = "This Req ID already exists.", True
rq.add_data_validation(v); v.add(f"A2:A{RL}")
v = DataValidation(type="custom", formula1=f"OR(C2=\"\",COUNTIF($C$2:$C${CL},C2)=1)", allow_blank=True)
v.errorTitle, v.error, v.showErrorMessage = "Duplicate candidate", "This email is already in the tracker.", True
cd.add_data_validation(v); v.add(f"C2:C{CL}")

# ------------------------------------------------------ conditional formats
def fill(c):
    return PatternFill("solid", fgColor=c, bgColor=c)


RED, AMBER, GREEN, GREYF = "F8CBAD", "FFE699", "C6EFCE", "E7E6E6"
cd.conditional_formatting.add(f"Z2:Z{CL}", CellIsRule(operator="equal", formula=['"BREACH"'], fill=fill(RED), font=Font(bold=True, color="9C0006")))
cd.conditional_formatting.add(f"Z2:Z{CL}", CellIsRule(operator="equal", formula=['"On Track"'], fill=fill(GREEN)))
cd.conditional_formatting.add(f"M2:M{CL}", CellIsRule(operator="equal", formula=['"Hired"'], fill=fill(GREEN), font=Font(bold=True, color="006100")))
cd.conditional_formatting.add(f"M2:M{CL}", CellIsRule(operator="equal", formula=['"Lost"'], fill=fill(GREYF)))
cd.conditional_formatting.add(f"V2:V{CL}", FormulaRule(formula=['AND(ISNUMBER($V2),$V2>0.4)'], fill=fill(AMBER)))
cd.conditional_formatting.add(f"A2:AE{CL}", FormulaRule(formula=['$M2="Lost"'], font=Font(color="808080")))
cd.conditional_formatting.add(f"N2:N{CL}", FormulaRule(formula=['AND(OR($K2="Rejected",$K2="Withdrawn"),$N2="")'], fill=fill(AMBER)))
rq.conditional_formatting.add(f"S2:S{RL}", CellIsRule(operator="equal", formula=['"Overdue"'], fill=fill(RED), font=Font(bold=True, color="9C0006")))
rq.conditional_formatting.add(f"S2:S{RL}", CellIsRule(operator="equal", formula=['"Filled"'], fill=fill(GREEN)))
rq.conditional_formatting.add(f"S2:S{RL}", CellIsRule(operator="equal", formula=['"No Pipeline"'], fill=fill(AMBER)))
rq.conditional_formatting.add(f"G2:G{RL}", CellIsRule(operator="equal", formula=['"High"'], font=Font(bold=True, color="C00000")))
from openpyxl.formatting.rule import DataBarRule
rq.conditional_formatting.add(f"P2:P{RL}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))

# ------------------------------------------------------------- sample data
if BLANK:
    reqs = [("REQ-001", "QA Automation Engineer", "QA", "Hiring Manager Name", "Chennai", 2, "High",
             RECRUITERS[0], date(2026, 10, 1), date(2026, 11, 15))]
    cands = [("Example Candidate", "example@email.com", "9876543210", "REQ-001", "LinkedIn", RECRUITERS[0],
              date(2026, 10, 5), "Tech Round 1", date(2026, 10, 10), None, 4.5, 8.0, 11.0, 30, 4, None, None)]
else:
    random.seed(7)
    TODAY = date(2026, 10, 15)
    reqs = [
        ("REQ-101", "QA Automation Engineer (Python)", "QA", "S. Raman", "Chennai", 3, "High", RECRUITERS[0], date(2026, 6, 1), date(2026, 9, 30)),
        ("REQ-102", "Senior SDET", "QA", "M. Iyer", "Bangalore", 2, "High", RECRUITERS[0], date(2026, 7, 10), date(2026, 10, 31)),
        ("REQ-103", "Python Backend Developer", "Engineering", "A. Kumar", "Chennai", 2, "Medium", RECRUITERS[1], date(2026, 7, 20), date(2026, 11, 15)),
        ("REQ-104", "DevOps Engineer", "DevOps", "R. Nair", "Hyderabad", 1, "Medium", RECRUITERS[2], date(2026, 8, 5), date(2026, 11, 30)),
        ("REQ-105", "Data Engineer", "Data", "P. Das", "Pune", 2, "Low", RECRUITERS[1], date(2026, 8, 20), date(2026, 12, 15)),
        ("REQ-106", "Cloud Architect", "Cloud", "K. Menon", "Remote", 1, "High", RECRUITERS[2], date(2026, 9, 15), date(2026, 12, 31)),
    ]
    first = ["Arun", "Divya", "Ravi", "Sneha", "Vikram", "Priya", "Karthik", "Meera", "Rahul", "Anjali",
             "Suresh", "Kavya", "Manoj", "Deepa", "Ajay", "Nisha", "Gokul", "Lakshmi", "Hari", "Pooja",
             "Naveen", "Swathi", "Ramesh", "Ishita", "Vijay", "Revathi", "Sanjay", "Harini", "Prakash", "Aarthi",
             "Dinesh", "Keerthi", "Ganesh", "Bhavya", "Sathish", "Nandini", "Mohan", "Asha", "Senthil", "Ritu"]
    last = ["Kumar", "R", "S", "Menon", "Reddy", "Nair", "Iyer", "Shah", "Patel", "Das"]
    weights = [("Sourced", 5), ("Screening", 5), ("Tech Round 1", 5), ("Tech Round 2", 4), ("HR Round", 3),
               ("Offered", 3), ("Offer Accepted", 3), ("Joined", 5), ("Offer Declined", 2),
               ("Rejected", 8), ("Withdrawn", 2)]
    pool = [s for s, w in weights for _ in range(w)]
    order = {s[0]: s[1] for s in STAGES}
    cands = []
    for i, fn in enumerate(first):
        rqi = reqs[i % len(reqs)]
        stage = random.choice(pool)
        applied = max(rqi[8] + timedelta(days=random.randint(1, 40)), rqi[8])
        if stage == "Joined":
            applied = rqi[8] + timedelta(days=random.randint(1, 20))
        applied = min(applied, TODAY - timedelta(days=5))
        cur = round(random.uniform(4, 25), 1)
        exp_ctc = round(cur * random.uniform(1.15, 1.55), 1)
        notice = random.choice([15, 30, 60, 90])
        offer = join = None
        dropped = None
        if order[stage] >= 6:
            offer = min(applied + timedelta(days=random.randint(18, 35)), TODAY - timedelta(days=random.randint(1, 20)))
        if stage == "Joined":
            join = min(offer + timedelta(days=notice), TODAY - timedelta(days=1))
        if stage in ("Rejected", "Withdrawn"):
            dropped = random.choice(DROP_STAGES[:5])
        upd = join or offer or max(applied, TODAY - timedelta(days=random.randint(0, 7)))
        if stage == "Offer Accepted":
            upd = min(offer + timedelta(days=random.randint(1, 4)), TODAY)
        if stage in ("Rejected", "Withdrawn"):
            upd = min(applied + timedelta(days=random.randint(2, 20)), TODAY)
        rating = None if order[stage] < 3 and stage not in ("Rejected", "Withdrawn") else random.choice([2, 3, 3, 4, 4, 5])
        name = f"{fn} {random.choice(last)}"
        cands.append((name, f"{fn.lower()}.{i}@mail.com", f"9{random.randint(100000000, 999999999)}",
                      rqi[0], random.choice(SOURCES[:6]), rqi[7], applied, stage, upd, dropped,
                      round(random.uniform(1, 14), 1), cur, exp_ctc, notice, rating, offer, join))

for r, row in enumerate(reqs, 2):
    for c, v in enumerate(row, 1):
        rq.cell(row=r, column=c, value=v)
in_cols = [c for c, (h, k, w, f) in enumerate(CAND_COLS, 1) if k == "in"]
for r, row in enumerate(cands, 2):
    for c, v in zip(in_cols, row):
        cd.cell(row=r, column=c, value=v)

if BLANK:
    for ws in (rq, cd):
        ws["A2"].comment = Comment("Example row — overwrite or delete it, the formulas stay.", "Tracker")

# --------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMN", [2, 20, 12, 12, 12, 3, 20, 12, 12, 12, 12, 3, 14, 12]):
    db.column_dimensions[col].width = w
db["B2"] = "Recruitment Dashboard"
db["B2"].font = font(bold=True, size=20, color=NAVY)
db["B3"] = '="Auto-refreshed • Report as of "&TEXT(TODAY(),"dd-mmm-yyyy")'
db["B3"].font = font(italic=True, color="595959")

C = lambda col: f"Candidates!${col}$2:${col}${CL}"
R = lambda col: f"Requisitions!${col}$2:${col}${RL}"
KPIS = [
    ("Total Candidates", f'=COUNTIF({C("B")},"?*")', "0"),
    ("Active Pipeline", f'=COUNTIF({C("M")},"Active")', "0"),
    ("Offers Made", f'=COUNTIF({C("T")},"<>")', "0"),
    ("Joined", f'=COUNTIF({C("K")},"Joined")', "0"),
    ("Offer Acceptance %", f'=IFERROR((COUNTIF({C("K")},"Offer Accepted")+COUNTIF({C("K")},"Joined"))/(COUNTIF({C("K")},"Offer Accepted")+COUNTIF({C("K")},"Joined")+COUNTIF({C("K")},"Offer Declined")),0)', "0%"),
    ("Avg Time to Hire (d)", f'=IFERROR(AVERAGE({C("AD")}),0)', "0"),
    ("Avg Time to Offer (d)", f'=IFERROR(AVERAGE({C("AC")}),0)', "0"),
    ("SLA Breaches", f'=COUNTIF({C("Z")},"BREACH")', "0"),
    ("Open Requisitions", f'=COUNTIF({R("S")},"Open")+COUNTIF({R("S")},"Overdue")+COUNTIF({R("S")},"No Pipeline")', "0"),
    ("Openings Left", f'=SUM({R("O")})', "0"),
    ("Overdue Reqs", f'=COUNTIF({R("S")},"Overdue")', "0"),
    ("Avg Hike Asked", f'=IFERROR(AVERAGE({C("V")}),0)', "0%"),
]
tile_cols = ["B", "C", "D", "E"]
tile_fill = PatternFill("solid", fgColor=GREY)
for i, (label, f, fmt) in enumerate(KPIS):
    blk, pos = divmod(i, 4)
    # 4 tiles per row; tile = label row + value row; tile spans two columns pairs
    col = ["B", "D", "G", "I"][pos]
    col2 = ["C", "E", "H", "J"][pos]
    r = 5 + blk * 3
    db.merge_cells(f"{col}{r}:{col2}{r}")
    db.merge_cells(f"{col}{r+1}:{col2}{r+1}")
    db[f"{col}{r}"] = label
    db[f"{col}{r}"].font = font(size=9, color="595959", bold=True)
    db[f"{col}{r+1}"] = f
    db[f"{col}{r+1}"].number_format = fmt
    db[f"{col}{r+1}"].font = font(size=18, bold=True, color=TEAL if label != "SLA Breaches" and label != "Overdue Reqs" else "C00000")
    for rr in (r, r + 1):
        for cc in (col, col2):
            db[f"{cc}{rr}"].fill = tile_fill
            db[f"{cc}{rr}"].alignment = Alignment(horizontal="center", vertical="center")
    db.row_dimensions[r + 1].height = 30


def table_header(ws, row, col_letters, headers):
    for cl, h in zip(col_letters, headers):
        c = ws[f"{cl}{row}"]
        c.value = h
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def section(ws, cell, text):
    ws[cell] = text
    ws[cell].font = font(bold=True, size=12, color=NAVY)


# Funnel
section(db, "B15", "Hiring Funnel")
table_header(db, 16, "BCDE", ["Stage", "Reached", "% of Total", "Step Conv."])
funnel = [s for s in STAGES if s[3] != "Lost"]
for i, (st, order, *_ ) in enumerate(funnel):
    r = 17 + i
    db[f"B{r}"] = st
    db[f"C{r}"] = f'=COUNTIF({C("AE")},">="&INDEX(Lists!$B$4:$B${STAGE_LAST},MATCH(B{r},{ST_RNG},0)))'
    db[f"D{r}"] = f"=IFERROR(C{r}/$C$17,0)"
    db[f"E{r}"] = "" if i == 0 else f"=IFERROR(C{r}/C{r-1},0)"
    db[f"D{r}"].number_format = db[f"E{r}"].number_format = "0%"
FUN_END = 16 + len(funnel)

# Current stage distribution
section(db, "G15", "Current Stage Snapshot")
table_header(db, 16, "GHI", ["Stage", "Candidates", "SLA Breach"])
for i, (st, *_ ) in enumerate(STAGES):
    r = 17 + i
    db[f"G{r}"] = st
    db[f"H{r}"] = f'=COUNTIF({C("K")},G{r})'
    db[f"I{r}"] = f'=COUNTIFS({C("K")},G{r},{C("Z")},"BREACH")'
STG_END = 16 + len(STAGES)

# Source effectiveness
r0 = STG_END + 3
section(db, f"B{r0}", "Source Effectiveness")
table_header(db, r0 + 1, "BCDE", ["Source", "Candidates", "Hires", "Hire Rate"])
for i, s in enumerate(SOURCES):
    r = r0 + 2 + i
    db[f"B{r}"] = f"=Lists!G{4 + i}"
    db[f"C{r}"] = f'=IF(B{r}="","",COUNTIF({C("H")},B{r}))'
    db[f"D{r}"] = f'=IF(B{r}="","",COUNTIFS({C("H")},B{r},{C("K")},"Joined"))'
    db[f"E{r}"] = f'=IF(B{r}="","",IFERROR(D{r}/C{r},0))'
    db[f"E{r}"].number_format = "0%"
SRC_START, SRC_END = r0 + 2, r0 + 1 + len(SOURCES)

# Recruiter performance
section(db, f"G{r0}", "Recruiter Performance")
table_header(db, r0 + 1, "GHIJK", ["Recruiter", "Candidates", "Active", "Offers", "Joined"])
for i in range(len(RECRUITERS)):
    r = r0 + 2 + i
    db[f"G{r}"] = f"=Lists!H{4 + i}"
    db[f"H{r}"] = f'=IF(G{r}="","",COUNTIF({C("I")},G{r}))'
    db[f"I{r}"] = f'=IF(G{r}="","",COUNTIFS({C("I")},G{r},{C("M")},"Active"))'
    db[f"J{r}"] = f'=IF(G{r}="","",COUNTIFS({C("I")},G{r},{C("T")},"<>"))'
    db[f"K{r}"] = f'=IF(G{r}="","",COUNTIFS({C("I")},G{r},{C("K")},"Joined"))'

# Monthly trend (current year)
m0 = SRC_END + 3
section(db, f"B{m0}", "Monthly Trend (current year)")
table_header(db, m0 + 1, "BCDE", ["Month", "Applications", "Offers", "Joined"])
for m in range(12):
    r = m0 + 2 + m
    db[f"B{r}"] = f"=DATE(YEAR(TODAY()),{m + 1},1)"
    db[f"B{r}"].number_format = "mmm-yy"
    for cl, src in zip("CDE", ["J", "T", "U"]):
        db[f"{cl}{r}"] = f'=COUNTIFS({C(src)},">="&$B{r},{C(src)},"<"&EDATE($B{r},1))'
MON_END = m0 + 13

# Requisition health
section(db, f"G{m0}", "Requisition Health")
table_header(db, m0 + 1, "GHIJK", ["Req Status", "Count", "", "", ""])
for i, st in enumerate(["Open", "Overdue", "No Pipeline", "Filled"]):
    r = m0 + 2 + i
    db[f"G{r}"] = st
    db[f"H{r}"] = f'=COUNTIF({R("S")},G{r})'
for cl in "IJK":
    db[f"{cl}{m0+1}"].value = None
    db[f"{cl}{m0+1}"].fill = PatternFill()

# Styling for all dashboard tables
for row in db.iter_rows(min_row=17, max_row=MON_END, min_col=2, max_col=11):
    for c in row:
        if c.value is not None and c.row not in (r0 + 1, m0 + 1) and not (c.font and c.font.b and c.font.size == 12):
            c.font = font()
            c.border = BORDER

# Charts
ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Hiring Funnel"
ch.add_data(Reference(db, min_col=3, min_row=16, max_row=FUN_END), titles_from_data=True)
ch.set_categories(Reference(db, min_col=2, min_row=17, max_row=FUN_END))
ch.y_axis.title = None; ch.x_axis.scaling.orientation = "maxMin"; ch.legend = None
ch.height, ch.width = 8, 14
db.add_chart(ch, "M4")

pie = PieChart(); pie.title = "Candidates by Source"
pie.add_data(Reference(db, min_col=3, min_row=SRC_START - 1, max_row=SRC_END), titles_from_data=True)
pie.set_categories(Reference(db, min_col=2, min_row=SRC_START, max_row=SRC_END))
pie.height, pie.width = 8, 14
db.add_chart(pie, "M21")

ln = LineChart(); ln.title = "Monthly Trend"; ln.style = 12
ln.add_data(Reference(db, min_col=3, max_col=5, min_row=m0 + 1, max_row=MON_END), titles_from_data=True)
ln.set_categories(Reference(db, min_col=2, min_row=m0 + 2, max_row=MON_END))
ln.x_axis.number_format = "mmm"; ln.height, ln.width = 8, 14
db.add_chart(ln, f"M{m0}")

# --------------------------------------------------------------- How To Use
hw = wb.create_sheet("How To Use", 0)
hw.sheet_view.showGridLines = False
hw.column_dimensions["A"].width = 2
hw.column_dimensions["B"].width = 110
lines = [
    ("Automated Recruitment Tracker", font(bold=True, size=18, color=NAVY)),
    ("You only type in the cream / dark-navy-header columns. Everything in light-blue cells (teal headers) is calculated automatically.", font(size=11)),
    ("", None),
    ("1. Requisitions — add each open role (Req ID, Role, Openings, Open & Target dates). Candidate counts, fill %, days open and status (Open / Overdue / No Pipeline / Filled) update themselves.", font()),
    ("2. Candidates — add one row per candidate. Pick the Req ID from the dropdown; Role and Hiring Manager fill in automatically. Candidate ID is generated.", font()),
    ("3. When a candidate moves, change 'Current Stage' and set 'Stage Updated On' to today (Ctrl + ; ). The sheet recalculates days in stage, SLA status, next action, hike %, time-to-offer and time-to-hire.", font()),
    ("4. For Rejected / Withdrawn candidates, fill 'Dropped At Stage' (cell turns amber until you do) so the funnel stays accurate.", font()),
    ("5. Dashboard — KPI tiles, funnel conversion, source effectiveness, recruiter performance, monthly trend and requisition health. Nothing to edit; it refreshes on open.", font()),
    ("6. Lists — change stages, SLA days, next-action text, sources, recruiters, departments and locations. Every dropdown and rule follows these settings.", font()),
    ("", None),
    ("Automations built in", font(bold=True, size=12, color=NAVY)),
    ("• Auto Candidate ID, Role / Hiring Manager lookup, duplicate email & Req ID blocking, dropdowns and date validation on every input.", font()),
    ("• SLA engine: flags candidates sitting in a stage longer than its SLA (red 'BREACH'). Next Action tells the recruiter what to do.", font()),
    ("• Expected joining date = Offer date + notice period; hike % highlighted amber when above 40%.", font()),
    ("• Requisitions turn red when past target date and not filled; data bars show fill %.", font()),
    ("• Pre-wired for 500 candidates and 100 requisitions — just keep typing in the next empty row.", font()),
    ("", None),
    ("Notes / assumptions", font(bold=True, size=12, color=NAVY)),
    ("• CTC is in Lakhs Per Annum (LPA). SLA days per stage are defaults set on the Lists sheet — adjust to your process.", font()),
    ("• Days counts use TODAY(), so the workbook updates every time it is opened.", font()),
    ("• Offer Acceptance % = (Offer Accepted + Joined) ÷ (Offer Accepted + Joined + Offer Declined).", font()),
]
if not BLANK:
    lines.append(("• This copy contains DEMO data (6 reqs, 40 candidates). Use Recruitment_Tracker_Template.xlsx for a clean start.", font(color="C00000")))
for i, (t, f) in enumerate(lines, 2):
    hw[f"B{i}"] = t
    hw[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    if f:
        hw[f"B{i}"].font = f
for i, (t, f) in enumerate(lines, 2):
    if len(t) > 110:
        hw.row_dimensions[i].height = 28

# Tab colors & order
hw.sheet_properties.tabColor = "7F7F7F"
db.sheet_properties.tabColor = NAVY
rq.sheet_properties.tabColor = TEAL
cd.sheet_properties.tabColor = TEAL
ls.sheet_properties.tabColor = "BF8F00"
wb.active = 1
wb.save(OUT)
print("saved", OUT)
